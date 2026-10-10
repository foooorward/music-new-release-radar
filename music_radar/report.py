"""产出层：每日 CSV / Excel 表格 + 单文件网页看板。

看板是自包含的单个 HTML（数据内嵌），双击就能打开，也能直接丢到任何静态托管上。
"""

from __future__ import annotations

import csv
import html
import json
from pathlib import Path

CSV_COLUMNS = [
    ("rank", "综合排名"),
    ("score", "热度分"),
    ("title", "歌名"),
    ("artist", "歌手"),
    ("album", "专辑"),
    ("platforms_text", "收录平台"),
    ("best_rank", "最佳名次"),
    ("best_chart", "所在榜单"),
    ("publish_time", "发布时间"),
    ("is_new_release", "近期新发行"),
    ("in_new_chart", "入选新歌榜"),
    ("platform_count", "跨平台数"),
    ("prev_rank", "上次名次"),
    ("rank_delta", "名次变化"),
    ("artist_hot", "音乐人热度"),
    ("parts_text", "分项(名次/覆盖/歌手/爬升)"),
    ("url", "链接"),
]


def _flatten(item: dict) -> dict:
    row = dict(item)
    row["artist"] = "、".join(item["artists"])
    row["platforms_text"] = " / ".join(item["platforms"])
    row["is_new_release"] = "是" if item["is_new_release"] else "否"
    row["in_new_chart"] = "是" if item.get("in_new_chart") else "否"
    row["prev_rank"] = item["prev_rank"] or ""
    row["rank_delta"] = "" if item["rank_delta"] is None else item["rank_delta"]
    parts = item["parts"]
    row["parts_text"] = "名次%s/覆盖%s/歌手%s/爬升%s" % (
        "-" if parts["rank"] is None else parts["rank"],
        parts["cross"],
        parts["artist"],
        parts["momentum"],
    )
    primary_platform = item["platform_keys"][0] if item["platform_keys"] else ""
    row["url"] = item["urls"].get(primary_platform, "")
    row["platform_count"] = item["platform_count"]
    return row


def write_csv(scored: list[dict], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig 让 Excel 直接双击打开不乱码
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for _, label in CSV_COLUMNS])
        for item in scored:
            row = _flatten(item)
            writer.writerow([row.get(key, "") for key, _ in CSV_COLUMNS])
    return path


def write_xlsx(scored: list[dict], path: Path) -> tuple[Path | None, str]:
    """写了 .xlsx 就当加分项；没装 openpyxl 就安静跳过，不影响流程。"""
    try:
        from openpyxl import Workbook  # type: ignore
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None, "未安装 openpyxl，已跳过 Excel（CSV 一样能用 Excel 打开）"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "新歌热度榜"

    headers = [label for _, label in CSV_COLUMNS]
    sheet.append(headers)
    head_fill = PatternFill("solid", fgColor="DCE7FB")
    for column, _ in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in scored:
        row = _flatten(item)
        sheet.append([row.get(key, "") for key, _ in CSV_COLUMNS])

    widths = [8, 8, 28, 20, 24, 26, 9, 14, 12, 12, 12, 9, 9, 10, 11, 30, 38]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))
    return path, ""


def write_dashboard(
    scored: list[dict],
    artist_table: dict[str, dict],
    source_status: list[dict],
    meta: dict,
    path: Path,
) -> Path:
    payload = {
        "meta": meta,
        "rows": [_flatten(item) | {"_parts": item["parts"], "_charts": item["charts"],
                                   "_urls": item["urls"], "_platform_keys": item["platform_keys"]}
                 for item in scored],
        "artists": sorted(
            (
                {
                    "name": row["name"],
                    "songs": row["songs"],
                    "hot": round(row["hot"] * 100, 1),
                    "music_size": row.get("music_size"),
                    "platforms": row.get("platforms") or [],
                }
                for row in artist_table.values()
            ),
            key=lambda x: (-x["hot"], -x["songs"]),
        )[:20],
        "sources": source_status,
    }
    serialized = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    rendered = DASHBOARD_TEMPLATE.replace("__PAYLOAD__", serialized)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(rendered, encoding="utf-8")
    return path


def write_summary_text(scored: list[dict], source_status: list[dict], meta: dict, top_n: int = 10) -> str:
    lines = []
    lines.append("=" * 78)
    lines.append("%s  ·  %s" % (meta.get("title", "每日新歌雷达"), meta["date"]))
    lines.append("=" * 78)
    lines.append("数据源状态：")
    for item in source_status:
        if item.get("skipped"):
            mark = "跳过"
        elif item.get("ok"):
            mark = "成功"
        else:
            mark = "失败"
        lines.append(
            "  [%s] %-18s %s 条 %s"
            % (mark, item["label"], item["count"], ("— " + item["note"]) if item["note"] else "")
        )
    lines.append("")
    lines.append("今日热度 Top %d：" % top_n)
    lines.append("  %-4s %-6s %-30s %-18s %-12s %s" % ("#", "热度", "歌名", "歌手", "发布时间", "平台"))
    for item in scored[:top_n]:
        title = item["title"][:28]
        artist = ("、".join(item["artists"]))[:16]
        lines.append(
            "  %-4s %-6s %-30s %-18s %-12s %s"
            % (
                item["rank"],
                item["score"],
                title,
                artist,
                item["publish_time"] or "—",
                "/".join(item["platforms"]),
            )
        )
    lines.append("")
    return "\n".join(lines)


DASHBOARD_TEMPLATE = r"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>每日新歌雷达</title>
<style>
:root{
  --bg:#f4f6fa; --card:#ffffff; --line:#e5e9f2; --line2:#eef1f7;
  --text:#1b2130; --muted:#6b7688; --muted2:#98a1b2;
  --primary:#2f6df6; --primary-l:#e8effe; --ok:#1f9d63; --warn:#c98118; --bad:#d64545;
}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--text);
  font-family:system-ui,-apple-system,"Segoe UI","PingFang SC","Microsoft YaHei",sans-serif;
  font-size:14px;line-height:1.6}
a{color:var(--primary);text-decoration:none}
a:hover{text-decoration:underline}
.wrap{max-width:1440px;margin:0 auto;padding:20px 22px 60px}
header.top{display:flex;flex-wrap:wrap;gap:14px;align-items:flex-end;justify-content:space-between;
  padding:18px 20px;background:var(--card);border:1px solid var(--line);border-radius:14px}
h1{margin:0;font-size:19px;font-weight:600;letter-spacing:.2px}
.sub{margin:4px 0 0;color:var(--muted);font-size:12.5px}
.chips{display:flex;flex-wrap:wrap;gap:8px}
.chip{display:inline-flex;align-items:center;gap:6px;padding:4px 10px;border-radius:999px;
  font-size:12px;border:1px solid var(--line);background:var(--line2);color:var(--muted)}
.chip .dot{width:6px;height:6px;border-radius:50%;background:var(--muted2)}
.chip.ok{background:#e7f6ee;border-color:#c6e8d6;color:#1a7a4d}
.chip.ok .dot{background:var(--ok)}
.chip.bad{background:#fdeaea;border-color:#f6cfcf;color:#a83232}
.chip.bad .dot{background:var(--bad)}
.chip.skip{background:#f3f4f7}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin:14px 0}
.stat{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:13px 16px}
.stat b{display:block;font-size:23px;font-weight:600;letter-spacing:.3px}
.stat span{color:var(--muted);font-size:12.5px}
.toolbar{display:flex;flex-wrap:wrap;gap:10px;align-items:center;
  background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px 14px;margin-bottom:14px}
.toolbar input[type=search]{flex:1 1 200px;min-width:160px;padding:8px 12px;border:1px solid var(--line);
  border-radius:9px;font-size:13px;outline:none;background:#fff;color:var(--text)}
.toolbar input[type=search]:focus{border-color:var(--primary)}
.pf{display:inline-flex;align-items:center;gap:6px;padding:6px 11px;border:1px solid var(--line);
  border-radius:999px;font-size:12.5px;cursor:pointer;user-select:none;background:#fff;color:var(--muted)}
.pf.on{background:var(--primary-l);border-color:#c3d6fb;color:#1d4fc4;font-weight:500}
.pf input{display:none}
.switch{display:inline-flex;align-items:center;gap:7px;font-size:12.5px;color:var(--muted);cursor:pointer}
.switch select{padding:6px 9px;border:1px solid var(--line);border-radius:8px;font-size:12.5px;
  background:#fff;color:var(--text);outline:none;cursor:pointer}
.switch select:focus{border-color:var(--primary)}
.grid{display:grid;grid-template-columns:minmax(0,1fr) 310px;gap:14px;align-items:start}
@media(max-width:1080px){.grid{grid-template-columns:1fr}}
.card{background:var(--card);border:1px solid var(--line);border-radius:12px;overflow:hidden}
.card h2{margin:0;padding:13px 16px;font-size:14px;font-weight:600;border-bottom:1px solid var(--line2);
  display:flex;justify-content:space-between;align-items:center;gap:10px}
.card h2 em{font-style:normal;color:var(--muted2);font-size:12px;font-weight:400}
.tablewrap{overflow-x:auto}
table{width:100%;border-collapse:collapse;font-size:13px}
th,td{padding:9px 10px;text-align:left;border-bottom:1px solid var(--line2);white-space:nowrap}
th{font-weight:500;color:var(--muted);font-size:12.5px;background:#fafbfd;position:sticky;top:0;z-index:2}
th.sortable{cursor:pointer;user-select:none}
th.sortable:hover{color:var(--primary)}
th .arrow{font-size:10px;color:var(--primary)}
tbody tr:hover{background:#fafbfe}
td.title{white-space:normal;min-width:190px;max-width:300px}
td.title b{font-weight:500}
td.title small{display:block;color:var(--muted2);font-size:11.5px;margin-top:1px}
.rankcell{color:var(--muted2);font-size:12.5px;width:38px}
.scorebar{display:flex;align-items:center;gap:8px;min-width:104px}
.scorebar i{display:block;height:6px;border-radius:4px;background:var(--primary)}
.scorebar span{font-variant-numeric:tabular-nums;font-weight:600;font-size:13px}
.tier-a i{background:#d64545}
.tier-b i{background:#e08a2b}
.tier-c i{background:#2f6df6}
.tier-d i{background:#8b98ad}
.pill{display:inline-block;padding:2px 8px;border-radius:999px;font-size:11.5px;
  border:1px solid var(--line);background:#f7f9fc;color:var(--muted);margin-right:4px}
.pill.netease{background:#fdeaea;border-color:#f6d0d0;color:#b03a3a}
.pill.kugou{background:#e9f1fd;border-color:#cddffa;color:#2b62c4}
.pill.itunes{background:#f0eefc;border-color:#d8d3f6;color:#5b4bc4}
.delta{font-size:12px}
.delta.up{color:#d64545}
.delta.down{color:#1f9d63}
.newtag{display:inline-block;padding:1px 7px;border-radius:5px;font-size:11px;
  background:#e8effe;color:#1d4fc4;border:1px solid #cddffa}
.artistrow{display:flex;align-items:center;gap:10px;padding:9px 16px;border-bottom:1px solid var(--line2)}
.artistrow:last-child{border-bottom:none}
.artistrow .na{flex:1 1 auto;min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;font-size:13px}
.artistrow .bar{flex:0 0 84px;height:6px;border-radius:4px;background:#eef1f7;overflow:hidden}
.artistrow .bar i{display:block;height:100%;background:#2f6df6}
.artistrow .val{flex:0 0 40px;text-align:right;font-size:12.5px;color:var(--muted);font-variant-numeric:tabular-nums}
.artistrow .no{flex:0 0 20px;color:var(--muted2);font-size:12px}
.more{padding:13px;text-align:center}
.more button{padding:8px 18px;border:1px solid var(--line);background:#fff;border-radius:9px;
  cursor:pointer;font-size:13px;color:var(--text)}
.more button:hover{border-color:var(--primary);color:var(--primary)}
.empty{padding:40px 16px;text-align:center;color:var(--muted2)}
.notes{margin-top:16px;background:var(--card);border:1px solid var(--line);border-radius:12px;padding:16px 18px}
.notes h3{margin:0 0 8px;font-size:13.5px;font-weight:600}
.notes ul{margin:0;padding-left:18px;color:var(--muted);font-size:12.5px}
.notes li{margin:4px 0}
.notes .warn{background:#fff8ec;border:1px solid #f2e0c0;border-radius:9px;padding:9px 12px;
  color:#8a6420;font-size:12.5px;margin-top:10px}
footer{margin-top:18px;text-align:center;color:var(--muted2);font-size:12px}
</style>
</head>
<body>
<div class="wrap">
  <header class="top">
    <div>
      <h1 id="pageTitle">每日新歌雷达</h1>
      <p class="sub" id="pageSub"></p>
    </div>
    <div class="chips" id="sourceChips"></div>
  </header>

  <section class="stats" id="statsRow"></section>

  <section class="toolbar">
    <input type="search" id="q" placeholder="搜索歌名 / 歌手 / 专辑">
    <span id="platformFilters"></span>
    <label class="switch">发布时间
      <select id="timeRange">
        <option value="0" selected>全部</option>
        <option value="1">今天发行</option>
        <option value="3">近 3 天发行</option>
        <option value="7">近 7 天发行</option>
        <option value="14">近 14 天发行</option>
        <option value="30">近 30 天发行</option>
      </select>
    </label>
    <label class="switch">数据范围
      <select id="scope">
        <option value="all">全部上榜作品</option>
        <option value="newchart" selected>仅平台新歌榜收录</option>
      </select>
    </label>
    <label class="switch"><input type="checkbox" id="multiOnly"> 只看跨平台收录</label>
  </section>

  <div class="grid">
    <section class="card">
      <h2>综合热度榜 <em id="tableMeta"></em></h2>
      <div class="tablewrap">
        <table>
          <thead>
            <tr id="headRow"></tr>
          </thead>
          <tbody id="tbody"></tbody>
        </table>
      </div>
      <div class="more" id="moreBox"></div>
    </section>

    <aside class="card">
      <h2>音乐人热度 <em>本次采集</em></h2>
      <div id="artistList"></div>
    </aside>
  </div>

  <section class="notes">
    <h3>数据口径与已知限制</h3>
    <ul>
      <li><b>默认只显示「被平台新歌榜收录」的作品</b>（网易云新歌榜 + 酷狗新歌榜）。热歌榜、TOP500 会收录大量经典老歌，混在一起会盖住真正的新歌；想看全量就把「数据范围」切成「全部上榜作品」。</li>
      <li><b>发布时间不是每个平台都有。</b>网易云提供准确发行日期（含专辑发布日期兜底，实测覆盖率 100%），Apple 提供发行日期；酷狗接口不提供歌曲发行日期——它的「榜单发布日期」是整个榜单共用的同一个时间戳，不能当作歌曲发行日期使用。所以酷狗独家的歌这里会留空，改用「入选新歌榜」来识别它是新歌。</li>
      <li><b>热度分</b> = 榜单名次 45% + 跨平台覆盖 25% + 音乐人热度 15% + 榜单爬升 15%，权重可在 <code>config.json</code> 调整。</li>
      <li><b>榜单名次</b>取该曲在所有真实榜单中按榜单权重加权平均后的名次分；Apple 的清单按发行时间排序，名次不代表热度，不参与计分。</li>
      <li><b>音乐人热度</b>是代理指标：本次榜单曝光度 + 作品体量（取自平台公开的作品数/专辑数）。所有国内平台都不公开歌手粉丝数，因此这里不是真实粉丝量。</li>
      <li><b>榜单爬升</b>需要至少两次采集才有意义。首次运行为中性值，第二天起开始显现真实变化。</li>
      <li>跨平台匹配按「主歌手 + 歌名」归一化比对，同名不同版本（Live、Remix）可能被合并或漏配。</li>
    </ul>
    <div class="warn">
      <b>关于播放量：</b>网易云、酷狗等国内平台均不对外公开真实播放量，其接口中的热度字段实测为空值，页面显示的"热度"是平台内部算法值。真正公开播放量数据的只有 YouTube（播放次数）与 Last.fm（播放次数），代码已预留接入位。因此本表按综合热度分排序，而不是按播放量排序。
    </div>
  </section>

  <footer id="collectedAt"></footer>
</div>

<script>
const DATA = __PAYLOAD__;
const PLATFORM_ORDER = ['netease', 'kugou', 'itunes'];
const PLATFORM_NAME = { netease: '网易云音乐', kugou: '酷狗音乐', itunes: 'Apple Music' };
const COLUMNS = [
  { key: 'rank',        label: '#',        sortable: true,  cls: 'rankcell' },
  { key: 'score',       label: '热度分',   sortable: true },
  { key: 'title',       label: '歌曲',     sortable: true,  cls: 'title' },
  { key: 'artist',      label: '歌手',     sortable: true },
  { key: 'platforms_text', label: '平台',  sortable: false },
  { key: 'best_rank',   label: '最佳名次', sortable: true },
  { key: 'best_chart',  label: '榜单',     sortable: false },
  { key: 'publish_time', label: '发布时间', sortable: true },
  { key: 'platform_count', label: '跨平台', sortable: true },
  { key: 'artist_hot',  label: '音乐人热度', sortable: true }
];

const state = {
  q: '',
  platforms: new Set(DATA.rows.flatMap(r => r._platform_keys).filter((v, i, a) => a.indexOf(v) === i)),
  days: 0,
  scope: 'newchart',
  multiOnly: false,
  sort: 'score',
  dir: -1,
  limit: 60
};

function esc(s) {
  return String(s === null || s === undefined ? '' : s)
    .replace(/&/g, '&amp;').replace(/</g, '&lt;').replace(/>/g, '&gt;')
    .replace(/"/g, '&quot;').replace(/'/g, '&#39;');
}

function tier(score) {
  if (score >= 80) return 'tier-a';
  if (score >= 65) return 'tier-b';
  if (score >= 45) return 'tier-c';
  return 'tier-d';
}

function fmtDuration(ms) {
  if (!ms) return '';
  const total = Math.round(ms / 1000);
  return Math.floor(total / 60) + ':' + String(total % 60).padStart(2, '0');
}

document.getElementById('pageTitle').textContent = DATA.meta.title;
document.getElementById('pageSub').textContent =
  '采集日期 ' + DATA.meta.date + ' · 收录 ' + DATA.meta.total + ' 首 · 去重后 ' + DATA.rows.length + ' 条作品记录';
document.getElementById('collectedAt').textContent = '采集时间 ' + DATA.meta.collected_at + ' · 数据来源为各平台公开页面与公开接口';

document.getElementById('sourceChips').innerHTML = DATA.sources.map(s => {
  const cls = s.skipped ? 'skip' : (s.ok ? 'ok' : 'bad');
  const label = s.skipped ? '已关闭' : (s.ok ? s.count + ' 条' : '失败');
  const title = s.note ? ' title="' + esc(s.note) + '"' : '';
  return '<span class="chip ' + cls + '"' + title + '><i class="dot"></i>' + esc(s.label) + ' · ' + label + '</span>';
}).join('');

const withDate = DATA.rows.filter(r => r.publish_time).length;
const stats = [
  { n: DATA.rows.length, t: '去重后作品数' },
  { n: DATA.rows.filter(r => r.is_new_release === '是').length, t: '近 14 天新发行' },
  { n: DATA.rows.filter(r => r.platform_count > 1).length, t: '跨平台收录' },
  { n: Math.round(withDate / Math.max(DATA.rows.length, 1) * 100) + '%', t: '发布时间覆盖率' },
  { n: DATA.sources.filter(s => s.ok).length + '/' + DATA.sources.length, t: '数据源可用' }
];
document.getElementById('statsRow').innerHTML = stats.map(s =>
  '<div class="stat"><b>' + s.n + '</b><span>' + s.t + '</span></div>').join('');

const platformFilters = document.getElementById('platformFilters');
platformFilters.innerHTML = PLATFORM_ORDER.filter(p => DATA.rows.some(r => r._platform_keys.includes(p))).map(p =>
  '<label class="pf on" data-p="' + p + '"><input type="checkbox" checked>' + PLATFORM_NAME[p] + '</label>'
).join('');
platformFilters.addEventListener('click', e => {
  const label = e.target.closest('.pf');
  if (!label) return;
  e.preventDefault();
  const p = label.dataset.p;
  if (state.platforms.has(p)) { state.platforms.delete(p); label.classList.remove('on'); }
  else { state.platforms.add(p); label.classList.add('on'); }
  state.limit = 60;
  render();
});

function buildHead() {
  document.getElementById('headRow').innerHTML = COLUMNS.map(c => {
    const active = state.sort === c.key;
    const arrow = active ? ' <span class="arrow">' + (state.dir === -1 ? '▼' : '▲') + '</span>' : '';
    return '<th class="' + (c.sortable ? 'sortable' : '') + '" data-key="' + c.key + '">' + c.label + arrow + '</th>';
  }).join('');
}
document.getElementById('headRow').addEventListener('click', e => {
  const th = e.target.closest('th');
  if (!th || !th.dataset.key) return;
  const key = th.dataset.key;
  if (!COLUMNS.find(c => c.key === key && c.sortable)) return;
  if (state.sort === key) state.dir *= -1;
  else { state.sort = key; state.dir = -1; }
  render();
});

function filtered() {
  let rows = DATA.rows.slice();
  const q = state.q.trim().toLowerCase();
  if (q) rows = rows.filter(r =>
    (r.title + ' ' + r.artist + ' ' + (r.album || '')).toLowerCase().includes(q));
  rows = rows.filter(r => r._platform_keys.some(p => state.platforms.has(p)));
  if (state.scope === 'newchart') rows = rows.filter(r => r.in_new_chart === '是');
  if (state.days > 0) {
    const base = new Date(DATA.meta.date + 'T00:00:00');
    const cutoff = new Date(base.getTime());
    cutoff.setDate(cutoff.getDate() - (state.days - 1));
    rows = rows.filter(r => {
      if (!r.publish_time) return false;
      const d = new Date(r.publish_time + 'T00:00:00');
      return !isNaN(d.getTime()) && d >= cutoff && d <= base;
    });
  }
  if (state.multiOnly) rows = rows.filter(r => r.platform_count > 1);
  const key = state.sort;
  rows.sort((a, b) => {
    let va = a[key], vb = b[key];
    if (key === 'publish_time') { va = va || ''; vb = vb || ''; }
    if (typeof va === 'number' && typeof vb === 'number') return (va - vb) * state.dir;
    return String(va).localeCompare(String(vb), 'zh') * state.dir;
  });
  return rows;
}

function render() {
  const rows = filtered();
  const shown = rows.slice(0, state.limit);
  document.getElementById('tableMeta').textContent =
    '显示 ' + shown.length + ' / ' + rows.length + ' 条';

  const tbody = document.getElementById('tbody');
  if (!shown.length) {
    tbody.innerHTML = '<tr><td colspan="' + COLUMNS.length + '"><div class="empty">没有符合条件的作品<br><span style="font-size:12px">试着把「发布时间」调成「全部」，或清空搜索词与平台筛选</span></div></td></tr>';
  } else {
    tbody.innerHTML = shown.map(r => {
      const bars = Math.max(3, Math.round(r.score));
      const scoreCell = '<div class="scorebar ' + tier(r.score) + '"><i style="width:' + bars + 'px"></i><span>' + r.score + '</span></div>';
      const pills = (r._platform_keys || []).map(p =>
        '<span class="pill ' + p + '">' + PLATFORM_NAME[p] + '</span>').join('');
      const link = r.url ? '<a href="' + esc(r.url) + '" target="_blank" rel="noopener">' + esc(r.title) + '</a>' : esc(r.title);
      const sub = (r.album ? esc(r.album) : '') + (r.publish_time ? (r.album ? ' · ' : '') + esc(r.publish_time) : '');
      let delta = '<span style="color:#98a1b2">—</span>';
      if (r.rank_delta !== '' && r.rank_delta !== null && r.rank_delta !== undefined) {
        if (r.rank_delta > 0) delta = '<span class="delta up">↑' + r.rank_delta + '</span>';
        else if (r.rank_delta < 0) delta = '<span class="delta down">↓' + Math.abs(r.rank_delta) + '</span>';
        else delta = '<span class="delta">持平</span>';
      } else if (r.prev_rank) {
        delta = '<span class="delta">上次 #' + r.prev_rank + '</span>';
      }
      const newTag = r.is_new_release === '是' ? ' <span class="newtag">新</span>' : '';
      return '<tr>' +
        '<td class="rankcell">' + r.rank + '</td>' +
        '<td>' + scoreCell + '</td>' +
        '<td class="title">' + link + newTag + (sub ? '<small>' + sub + '</small>' : '') + '</td>' +
        '<td>' + esc(r.artist) + '</td>' +
        '<td>' + pills + '</td>' +
        '<td>' + (r.best_rank || '—') + '</td>' +
        '<td>' + esc(r.best_chart || '—') + '</td>' +
        '<td>' + esc(r.publish_time || '—') + '</td>' +
        '<td>' + r.platform_count + ' 平台</td>' +
        '<td>' + r.artist_hot + '</td>' +
        '</tr>';
    }).join('');
  }

  const more = document.getElementById('moreBox');
  more.innerHTML = rows.length > state.limit
    ? '<button id="moreBtn">加载更多（还有 ' + (rows.length - state.limit) + ' 条）</button>'
    : '';
  const btn = document.getElementById('moreBtn');
  if (btn) btn.addEventListener('click', () => { state.limit += 60; render(); });
}

function renderArtists() {
  const list = document.getElementById('artistList');
  if (!DATA.artists.length) {
    list.innerHTML = '<div class="empty">暂无数据</div>';
    return;
  }
  const max = DATA.artists[0].hot || 1;
  list.innerHTML = DATA.artists.map((a, i) =>
    '<div class="artistrow">' +
      '<span class="no">' + (i + 1) + '</span>' +
      '<span class="na" title="' + esc(a.name) + '">' + esc(a.name) +
        (a.music_size ? '<span style="color:#98a1b2;font-size:11.5px"> · ' + a.songs + '首上榜/' + a.music_size + '作品</span>'
                      : '<span style="color:#98a1b2;font-size:11.5px"> · ' + a.songs + '首上榜</span>') +
      '</span>' +
      '<span class="bar"><i style="width:' + Math.round(a.hot / max * 100) + '%"></i></span>' +
      '<span class="val">' + a.hot + '</span>' +
    '</div>').join('');
}

document.getElementById('q').addEventListener('input', e => {
  state.q = e.target.value;
  state.limit = 60;
  render();
});
document.getElementById('timeRange').addEventListener('change', e => {
  state.days = parseInt(e.target.value, 10) || 0;
  state.limit = 60;
  render();
});
document.getElementById('scope').addEventListener('change', e => {
  state.scope = e.target.value;
  state.limit = 60;
  render();
});
document.getElementById('multiOnly').addEventListener('change', e => {
  state.multiOnly = e.target.checked;
  state.limit = 60;
  render();
});

buildHead();
render();
renderArtists();
</script>
</body>
</html>
"""


def html_escape(text: str) -> str:
    return html.escape(text or "")
